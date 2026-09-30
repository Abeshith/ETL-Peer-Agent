import csv
import json
import re
from pathlib import Path
from typing import List, Optional, Dict, Any
from datetime import datetime
from utils.models import TestResult, CodeAnalysis, BatchRunRow
from src.config import Config

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


class ReportGenerator:
    def __init__(self, analysis: CodeAnalysis):
        self.analysis = analysis
        self.output_dir = Config.OUTPUTS_DIR
        self.output_dir.mkdir(exist_ok=True)

    def generate_csv(self, filename: Optional[str] = None) -> Path:
        if not filename:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"etl_validation_report_{timestamp}.csv"

        filepath = self.output_dir / filename
        fieldnames = ["Test_ID", "Test_Name", "Phase", "Status", "Finding", "Recommendation", "Timestamp"]

        rows = []
        for result in self.analysis.test_results:
            if result.status == "GENERATED":
                continue
            if result.phase in ("Test Generation", "Execution Summary", "AI Analysis"):
                continue
            rows.append((result, result.phase, result.status, result.finding, result.recommendation))

        with open(filepath, "w", newline="", encoding="utf-8") as csvfile:
            writer = csv.DictWriter(csvfile, fieldnames=fieldnames)
            writer.writeheader()
            for seq, (r, phase, status, finding, recommendation) in enumerate(rows, start=1):
                # Keep raw finding without markdown formatting for CSV
                writer.writerow({
                    "Test_ID": f"TC{seq:03d}",
                    "Test_Name": r.test_name,
                    "Phase": phase,
                    "Status": status,
                    "Finding": finding,
                    "Recommendation": recommendation,
                    "Timestamp": r.timestamp.isoformat(),
                })
        return filepath

    def generate_excel(self, filename: Optional[str] = None) -> Optional[Path]:
        """Generate Excel report with actual bold formatting for finding keys."""
        if not HAS_OPENPYXL:
            print("[WARNING] openpyxl not installed; skipping Excel report generation")
            return None

        if not filename:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"etl_validation_report_{timestamp}.xlsx"

        filepath = self.output_dir / filename
        wb = Workbook()
        ws = wb.active
        ws.title = "Validation Report"

        # Header row with formatting
        headers = ["Test_ID", "Test_Name", "Phase", "Status", "Finding", "Recommendation", "Timestamp"]
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        
        for col_num, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Collect rows
        rows = []
        for result in self.analysis.test_results:
            if result.status == "GENERATED":
                continue
            if result.phase in ("Test Generation", "Execution Summary", "AI Analysis"):
                continue
            rows.append((result, result.phase, result.status, result.finding, result.recommendation))

        # Add data rows with bold keys in findings
        for seq, (r, phase, status, finding, recommendation) in enumerate(rows, start=2):
            ws.cell(row=seq, column=1).value = f"TC{seq-1:03d}"
            ws.cell(row=seq, column=2).value = r.test_name
            ws.cell(row=seq, column=3).value = phase
            ws.cell(row=seq, column=4).value = status
            
            # Add finding with bold keys
            finding_cell = ws.cell(row=seq, column=5)
            self._set_finding_with_bold_keys(finding_cell, finding)
            
            ws.cell(row=seq, column=6).value = recommendation
            ws.cell(row=seq, column=7).value = r.timestamp.isoformat()
            
            # Alignment and wrapping
            for col in [1, 2, 3, 4, 5, 6, 7]:
                cell = ws.cell(row=seq, column=col)
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

        # Adjust column widths
        ws.column_dimensions["A"].width = 10  # Test_ID
        ws.column_dimensions["B"].width = 35  # Test_Name
        ws.column_dimensions["C"].width = 25  # Phase
        ws.column_dimensions["D"].width = 12  # Status
        ws.column_dimensions["E"].width = 60  # Finding
        ws.column_dimensions["F"].width = 40  # Recommendation
        ws.column_dimensions["G"].width = 25  # Timestamp

        # Set row height for header
        ws.row_dimensions[1].height = 25

        wb.save(filepath)
        return filepath

    def _set_finding_with_bold_keys(self, cell, finding: str):
        """Set cell value with bold keys parsed from finding string using rich text."""
        if not finding or not isinstance(finding, str):
            cell.value = finding
            return

        if not HAS_OPENPYXL:
            cell.value = finding
            return

        # Use rich text to format bold keys with InlineFont
        parts = []
        # Split by pipe to handle each key:value pair
        segments = finding.split("|")
        
        for i, segment in enumerate(segments):
            segment = segment.strip()
            
            if not segment:
                continue
            
            # Match "key: value" pattern
            match = re.match(r'^([a-zA-Z_][a-zA-Z0-9_\-]*)\s*:\s*(.*)$', segment)
            
            if match:
                key, value = match.groups()
                # InlineFont for rich text (bold for key)
                bold_font = InlineFont(b=True)
                normal_font = InlineFont(b=False)
                
                # Add key in bold
                parts.append(TextBlock(bold_font, f"{key}: "))
                # Add value in normal
                parts.append(TextBlock(normal_font, value))
                
                # Add separator for next part (except last)
                if i < len(segments) - 1:
                    parts.append(TextBlock(normal_font, " | "))
            else:
                # No key:value pattern, add as regular text
                normal_font = InlineFont(b=False)
                parts.append(TextBlock(normal_font, segment))
                
                if i < len(segments) - 1:
                    parts.append(TextBlock(normal_font, " | "))
        
        if parts:
            try:
                cell.value = CellRichText(*parts)
            except Exception:
                # Fallback if rich text fails
                cell.value = finding
        else:
            cell.value = finding

    def generate_json(self, filename: Optional[str] = None) -> Path:
        if not filename:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"etl_analysis_{timestamp}.json"

        filepath = self.output_dir / filename

        skip_phases = {"Phase5C", "Phase6"}
        phase5b_by_id = {
            r.test_id: r
            for r in self.analysis.test_results
            if r.phase == "Phase5B"
        }

        test_results_out = []
        seq = 1
        for result in self.analysis.test_results:
            if result.phase in skip_phases:
                continue
            if result.phase == "Phase5B":
                continue
            phase_label = result.phase
            status = result.status
            finding = result.finding
            recommendation = result.recommendation
            if result.phase == "Phase5A" and result.status == "GENERATED":
                executed = phase5b_by_id.get(result.test_id)
                if executed:
                    phase_label = "Test Cases Execution"
                    status = executed.status
                    finding = executed.finding
                    recommendation = executed.recommendation
                else:
                    continue  # no execution result — drop from output entirely
            
            test_results_out.append({
                "test_id": f"TC{seq:03d}",
                "test_name": result.test_name,
                "phase": phase_label,
                "status": status,
                "finding": finding,
                "recommendation": recommendation,
                "timestamp": result.timestamp.isoformat(),
            })
            seq += 1

        phases_summary = {}
        for phase in ["Metadata Extraction", "Azure Validation", "Snowflake Validation",
                      "Resource Validation", "Pre-Execution Summary", "Post-Execution Report",
                      "Test Cases Execution", "AI Analysis"]:
            phase_rows = [r for r in test_results_out if r["phase"] == phase]
            if phase_rows:
                phases_summary[phase] = {
                    "total": len(phase_rows),
                    "passed": sum(1 for r in phase_rows if r["status"] == "PASS"),
                    "failed": sum(1 for r in phase_rows if r["status"] == "FAIL"),
                    "skipped": sum(1 for r in phase_rows if r["status"] == "SKIPPED"),
                }

        output_data = {
            "analysis_timestamp": self.analysis.analysis_timestamp.isoformat(),
            "file_analyzed": self.analysis.file_path,
            "metadata": {
                "credential_mechanism": self._sanitize_connection(self.analysis.metadata.azure_connection),
                "source_identifier": self.analysis.metadata.azure_container,
                "snowflake_account": self._mask_account(self.analysis.metadata.snowflake_connection),
                "target_database": self.analysis.metadata.target_database,
                "target_schema": self.analysis.metadata.target_schema,
                "target_table": self.analysis.metadata.target_table,
                "audit_table": self.analysis.metadata.audit_table,
            },
            "phases": phases_summary,
            "test_results": test_results_out,
        }

        with open(filepath, "w", encoding="utf-8") as f:
            json.dump(output_data, f, indent=2)

        return filepath

    def _sanitize_connection(self, raw: str) -> str:
        """Convert raw Python call expression to a readable credential mechanism label."""
        if not raw:
            return "Unknown"
        if "AyAzureKeyVault" in raw or "azureKV" in raw:
            return "Azure Key Vault (AyAzureKeyVault)"
        if "SecretClient" in raw or "azure.keyvault" in raw:
            return "Azure Key Vault (SecretClient)"
        if "secretsmanager" in raw or "get_secret_value" in raw:
            return "AWS Secrets Manager"
        if "hvac" in raw or "vault.read" in raw:
            return "HashiCorp Vault"
        if "os.environ" in raw or "os.getenv" in raw:
            return "Environment Variables"
        if "yaml.safe_load" in raw or "config" in raw.lower():
            return "Config File"
        # Already a clean label (not a Python expression)
        if "(" not in raw:
            return raw
        return raw

    def _mask_account(self, raw: str) -> str:
        """Keep role/user visible but mask the account identifier."""
        if not raw:
            return "Not configured"
        import re
        # Replace account value with masked version
        masked = re.sub(r'account=([^,\s]+)', 'account=***', raw)
        return masked

    def generate_summary_report(self) -> str:
        pass_count = self.analysis.get_pass_count()
        fail_count = self.analysis.get_fail_count()
        gen_count = self.analysis.get_generated_count()
        skip_count = self.analysis.get_skipped_count()
        executed_count = self.analysis.get_executed_count()
        total_tests = len(self.analysis.test_results)

        # Phase 5C summary takes priority for execution counts
        s = self.analysis.phase5c_summary
        if s:
            executed_count = s.get("executed", executed_count)
            pass_count_exec = s.get("passed", pass_count)
            fail_count_exec = s.get("failed", fail_count)
            skipped_exec = s.get("skipped", skip_count)
            risk = s.get("risk", "UNKNOWN")
            duration = s.get("duration_str", "N/A")
        else:
            pass_count_exec = pass_count
            fail_count_exec = fail_count
            skipped_exec = skip_count
            risk = "HIGH" if fail_count >= 3 else ("MEDIUM" if fail_count > 0 else "LOW")
            duration = "N/A"

        report = f"""
{'='*70}
ETL TEST CASE AGENT - ANALYSIS REPORT
{'='*70}

File Analyzed: {self.analysis.file_path}
Analysis Timestamp: {self.analysis.analysis_timestamp.isoformat()}

METADATA EXTRACTED:
  Azure Container: {self.analysis.metadata.azure_container or 'N/A'}
  Target Database: {self.analysis.metadata.target_database or 'N/A'}
  Target Schema:   {self.analysis.metadata.target_schema or 'N/A'}
  Target Table:    {self.analysis.metadata.target_table or 'N/A'}
  Audit Table:     {self.analysis.metadata.audit_table or 'N/A'}

EXECUTION SUMMARY:
  Executed Tests : {executed_count}
  Passed         : {pass_count_exec}
  Failed         : {fail_count_exec}
  Skipped        : {skipped_exec}
  Generated      : {gen_count}
  Duration       : {duration}
  Risk           : {risk}
  Status         : {'READY FOR STAGE 2' if fail_count_exec == 0 else 'INCOMPLETE - Fix Issues'}

PHASE BREAKDOWN:
"""
        for phase in ["Phase1", "Phase2", "Phase3", "Phase4", "Phase5A", "Phase5B", "Phase5C", "Phase6"]:
            phase_results = self.analysis.get_results_by_phase(phase)
            if not phase_results:
                continue
            
            phase_pass = sum(1 for r in phase_results if r.status == "PASS")
            phase_fail = sum(1 for r in phase_results if r.status == "FAIL")
            phase_skip = sum(1 for r in phase_results if r.status == "SKIPPED")
            phase_gen  = sum(1 for r in phase_results if r.status == "GENERATED")
            
            parts = [f"{phase_pass} PASS", f"{phase_fail} FAIL", f"{phase_skip} SKIPPED"]
            if phase_gen:
                parts.append(f"{phase_gen} GENERATED")
            report += f"  {phase}: {' / '.join(parts)}\n"

        report += f"\n{'='*70}\n"
        return report
